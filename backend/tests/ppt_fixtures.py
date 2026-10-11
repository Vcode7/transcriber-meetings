"""Synthetic programme-review deck used by the PPT insight and matching tests."""
from datetime import date

MEETING_DATE = date(2026, 10, 12)

DECK = [
    ("Atlas Programme Review", ["Steering Committee", "October 2026"]),
    ("Agenda", ["Programme timeline", "Atlas build status", "Beacon mobile app", "Risks and decisions"]),
    ("Programme Timeline", None),  # table slide, see TIMELINE_ROWS
    ("Project Atlas – Build Status", [
        "SAP integration 70% complete; on track for Phase 3 exit in Dec 2026",
        "Data migration dry-run 2 delayed from 15-Sep-2026 to 30-Oct-2026 due to vendor data quality issues",
        "API gateway deployed to production on 05-Aug-2026 ✓",
        "Performance testing scheduled for 20-Nov-2026",
    ]),
    ("Project Beacon – Mobile App", [
        "Beta released to 500 pilot users on 20-Jul-2026",
        "App store submission planned for 15-Nov-2026",
        "Depends on Atlas API v2 being available by 01-Nov-2026",
        "Crash rate reduced from 4.1% to 1.2% since beta",
    ]),
    ("Programme History", [
        "Programme approved by the board in March 2019",
        "Initial budget of $12M allocated in 2020",
        "Vendor selection completed in 2021",
        "Legacy CRM decommissioned in 2022",
    ]),
    ("End of Section 1", ["Next: risks, decisions and actions"]),
    ("Risks, Decisions and Actions", [
        "Risk: attrition of two senior vendor engineers may slow integration testing",
        "Decision: steering committee approved additional QA budget of $250K on 01-Oct-2026",
        "Action: Priya to finalise the cutover plan by 31-Oct-2026",
        "Action: Data team to clean customer master data before dry-run 2",
    ]),
    ("Long-term Vision 2030", [
        "Expand the platform to five regions by 2030",
        "Target 10 million monthly active users in 2031",
    ]),
    ("Conclusion", [
        "Phase 3 is on track except for the data migration slip",
        "Focus for the next month: cutover readiness and Beacon store submission",
    ]),
    ("Thank You", ["Questions?"]),
    ("References", ["Gartner (2023) ERP migration benchmarks", "Internal audit report 2024"]),
    ("Appendix A – Detailed Budget", ["Phase 1: $2.1M", "Phase 2: $3.4M", "Phase 3: $4.8M"]),
]

TIMELINE_ROWS = [
    ("Phase", "Start", "End", "Status"),
    ("Phase 1 – Discovery", "Jan 2025", "Jun 2025", "Completed"),
    ("Phase 2 – Design", "Jul 2025", "Mar 2026", "Completed"),
    ("Phase 3 – Build & Integration", "Apr 2026", "Dec 2026", "In progress"),
    ("Phase 4 – UAT & Go-live", "Jan 2027", "Apr 2027", "Planned"),
]


def build_pptx(path: str) -> str:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    for title, bullets in DECK:
        slide = prs.slides.add_slide(layout if bullets is not None else prs.slide_layouts[5])
        slide.shapes.title.text = title
        if bullets is None:
            rows, cols = len(TIMELINE_ROWS), len(TIMELINE_ROWS[0])
            table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(9), Inches(3)).table
            for r, row in enumerate(TIMELINE_ROWS):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for b in bullets[1:]:
                body.add_paragraph().text = b
    prs.save(path)
    return path


LONG_POINTS = [
    {"id": "lp1", "text": (
        "The team discussed the Atlas data migration. Dry run two slipped because the vendor's customer data had "
        "quality problems, and it is now expected at the end of October. Priya said the cutover plan cannot be "
        "finalised until the dry run passes, and the data team agreed to clean the customer master data this week."
    )},
    {"id": "lp2", "text": (
        "For the Beacon mobile app, the store submission is still targeted for mid November. The mobile lead "
        "flagged that this depends on the Atlas API version two, and asked for confirmation that it will be ready "
        "by the first of November."
    )},
    {"id": "lp3", "text": (
        "The committee confirmed the extra QA budget of 250 thousand dollars and asked the vendor to backfill the "
        "two senior engineers who resigned so that integration testing is not slowed down."
    )},
    {"id": "lp4", "text": (
        "Facilities gave an update on the office move to the new building and the parking arrangements for staff "
        "from next quarter."
    )},
]
